from pathlib import Path
import csv
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptInches, Pt as PptPt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[10]
ARTIFACT_DIR = Path(__file__).resolve().parent
SOURCE_MEMO = ROOT / "skill-pair-redteam/fixtures/pack_a/run_002/source_memo.md"
PILOT_METRICS = ROOT / "skill-pair-redteam/fixtures/pack_a/run_002/pilot_metrics.csv"

DOCX_PATH = ARTIFACT_DIR / "pilot_program_review_memo.docx"
PDF_PATH = ARTIFACT_DIR / "pilot_program_review_memo.pdf"
PPTX_PATH = ARTIFACT_DIR / "pilot_program_summary_deck.pptx"
SUMMARY_MD_PATH = ARTIFACT_DIR / "structured_memo_summary.md"

THEME = {
    "blue": "0066FF",
    "cyan": "00FFFF",
    "dark": "1E1E1E",
    "white": "FFFFFF",
    "light": "F4F8FB",
    "muted": "556070",
}


def read_inputs():
    memo_text = SOURCE_MEMO.read_text(encoding="utf-8")
    metrics = []
    with PILOT_METRICS.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            row["baseline"] = float(row["baseline"])
            row["pilot"] = float(row["pilot"])
            metrics.append(row)
    return memo_text, metrics


def extract_bullets(section_name, text):
    pattern = rf"## {re.escape(section_name)}\n(?P<body>.*?)(?:\n## |\Z)"
    match = re.search(pattern, text, flags=re.S)
    if not match:
        return []
    bullets = []
    for line in match.group("body").splitlines():
        line = line.strip()
        if line.startswith("- "):
            bullets.append(line[2:])
    return bullets


def metric_line(row):
    name = row["metric"].replace("_", " ").title()
    baseline = row["baseline"]
    pilot = row["pilot"]
    unit = row["unit"]
    if row["metric"] in {"turnaround_time", "formatting_rework"}:
        change = baseline - pilot
        return f"{name}: improved from {baseline:g} to {pilot:g} {unit}, a reduction of {change:g} {unit}."
    change = pilot - baseline
    return f"{name}: increased from {baseline:g} to {pilot:g} {unit}, a gain of {change:g} {unit}."


def build_summary(memo_text, metrics):
    main_points = extract_bullets("Main Points", memo_text)
    recommended = extract_bullets("Recommended Message", memo_text)
    if not recommended:
        match = re.search(r"## Recommended Message\n(?P<body>.*)", memo_text, flags=re.S)
        recommended = [match.group("body").strip()] if match else []
    metric_points = [metric_line(row) for row in metrics]
    recommendation = recommended[0] if recommended else (
        "Continue the pilot for one more quarter with clearer review ownership."
    )
    return {
        "title": "Document Automation Pilot Review",
        "audience": "Executive operations and document workflow stakeholders",
        "main_points": main_points,
        "metric_points": metric_points,
        "recommendation": recommendation,
        "risks": [
            "Source data quality remains inconsistent across intake materials.",
            "Final review ownership needs a clearer single accountable role.",
        ],
        "actions": [
            "Extend the pilot for one additional quarter.",
            "Add a reviewer checklist to reduce ambiguity before final review.",
            "Publish a shared theme package for memo and deck outputs.",
        ],
    }


def write_summary_md(summary):
    lines = [
        f"# {summary['title']}",
        "",
        f"Audience: {summary['audience']}",
        "",
        "## Executive Message",
        summary["recommendation"],
        "",
        "## Main Points",
    ]
    lines.extend(f"- {item}" for item in summary["main_points"])
    lines.extend(["", "## Pilot Metrics"])
    lines.extend(f"- {item}" for item in summary["metric_points"])
    lines.extend(["", "## Risks"])
    lines.extend(f"- {item}" for item in summary["risks"])
    lines.extend(["", "## Recommended Actions"])
    lines.extend(f"- {item}" for item in summary["actions"])
    SUMMARY_MD_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def make_docx(summary, metrics):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    for style_name, size in [("Heading 1", 18), ("Heading 2", 13)]:
        styles[style_name].font.name = "Arial"
        styles[style_name].font.size = Pt(size)
        styles[style_name].font.bold = True
        styles[style_name].font.color.rgb = RGBColor(0x1E, 0x1E, 0x1E)

    title = doc.add_heading(summary["title"], level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("Structured executive memo based on source memo and pilot metrics")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.color.rgb = RGBColor(0x55, 0x60, 0x70)

    doc.add_heading("Executive Message", level=2)
    doc.add_paragraph(summary["recommendation"])

    doc.add_heading("Main Points", level=2)
    for item in summary["main_points"]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Pilot Metrics", level=2)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for idx, label in enumerate(["Metric", "Baseline", "Pilot", "Change"]):
        cell = table.rows[0].cells[idx]
        cell.text = label
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row in metrics:
        cells = table.add_row().cells
        cells[0].text = row["metric"].replace("_", " ").title()
        cells[1].text = f"{row['baseline']:g} {row['unit']}"
        cells[2].text = f"{row['pilot']:g} {row['unit']}"
        delta = row["pilot"] - row["baseline"]
        cells[3].text = f"{delta:+g} {row['unit']}"

    doc.add_heading("Risks", level=2)
    for item in summary["risks"]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Recommended Actions", level=2)
    for item in summary["actions"]:
        doc.add_paragraph(item, style="List Number")

    doc.save(DOCX_PATH)


def make_pdf(summary, metrics):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="CenterTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        textColor=colors.HexColor("#1e1e1e"),
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
    ))
    styles.add(ParagraphStyle(
        name="Section",
        parent=styles["Heading2"],
        textColor=colors.HexColor("#0066ff"),
        fontName="Helvetica-Bold",
        fontSize=12,
        spaceBefore=12,
        spaceAfter=6,
    ))
    styles["Normal"].fontName = "Helvetica"
    styles["Normal"].fontSize = 9.5
    styles["Normal"].leading = 12

    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
    )
    story = [
        Paragraph(summary["title"], styles["CenterTitle"]),
        Paragraph("Structured executive memo based on source memo and pilot metrics", styles["Normal"]),
        Spacer(1, 0.15 * inch),
        Paragraph("Executive Message", styles["Section"]),
        Paragraph(summary["recommendation"], styles["Normal"]),
        Paragraph("Main Points", styles["Section"]),
    ]
    for item in summary["main_points"]:
        story.append(Paragraph(f"- {item}", styles["Normal"]))

    story.append(Paragraph("Pilot Metrics", styles["Section"]))
    data = [["Metric", "Baseline", "Pilot", "Change"]]
    for row in metrics:
        delta = row["pilot"] - row["baseline"]
        data.append([
            row["metric"].replace("_", " ").title(),
            f"{row['baseline']:g} {row['unit']}",
            f"{row['pilot']:g} {row['unit']}",
            f"{delta:+g} {row['unit']}",
        ])
    table = Table(data, colWidths=[2.4 * inch, 1.3 * inch, 1.3 * inch, 1.2 * inch])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e1e1e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cdd6e0")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f8fb")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(table)
    story.append(Paragraph("Risks", styles["Section"]))
    for item in summary["risks"]:
        story.append(Paragraph(f"- {item}", styles["Normal"]))
    story.append(Paragraph("Recommended Actions", styles["Section"]))
    for idx, item in enumerate(summary["actions"], start=1):
        story.append(Paragraph(f"{idx}. {item}", styles["Normal"]))
    doc.build(story)


def add_title(slide, title, dark=False):
    color = THEME["white"] if dark else THEME["dark"]
    box = slide.shapes.add_textbox(PptInches(0.55), PptInches(0.3), PptInches(8.9), PptInches(0.55))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = "DejaVu Sans"
    run.font.bold = True
    run.font.size = PptPt(25)
    run.font.color.rgb = PptRGBColor.from_string(color)


def add_body_text(slide, items, x, y, w, h, font_size=13, color=None):
    box = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = PptPt(6)
        p.font.name = "DejaVu Sans"
        p.font.size = PptPt(font_size)
        p.font.color.rgb = PptRGBColor.from_string(color or THEME["dark"])
    return box


def make_pptx(summary, metrics):
    prs = Presentation()
    prs.slide_width = PptInches(10)
    prs.slide_height = PptInches(5.625)
    blank = prs.slide_layouts[6]

    # Slide 1
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PptRGBColor.from_string(THEME["dark"])
    slide.shapes.add_shape(1, PptInches(0), PptInches(4.75), PptInches(10), PptInches(0.9)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = PptRGBColor.from_string(THEME["blue"])
    add_title(slide, summary["title"], dark=True)
    add_body_text(slide, [
        "Pilot review for document automation workflow",
        "Turnaround time fell from 5 days to 2 days while processed volume rose from 18 to 42 documents.",
    ], 0.7, 1.45, 6.2, 1.4, 18, THEME["white"])
    for idx, stat in enumerate([("60%", "faster turnaround"), ("2.3x", "documents processed"), ("4.5", "reviewer satisfaction")]):
        x = 0.75 + idx * 3.05
        slide.shapes.add_shape(1, PptInches(x), PptInches(3.55), PptInches(2.35), PptInches(0.85)).fill.solid()
        slide.shapes[-1].fill.fore_color.rgb = PptRGBColor.from_string(THEME["white"])
        add_body_text(slide, [stat[0], stat[1]], x + 0.12, 3.62, 2.1, 0.55, 11, THEME["dark"])

    # Slide 2
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PptRGBColor.from_string(THEME["light"])
    add_title(slide, "Main Points")
    for idx, item in enumerate(summary["main_points"]):
        x = 0.65 + (idx % 2) * 4.55
        y = 1.15 + (idx // 2) * 1.55
        shape = slide.shapes.add_shape(1, PptInches(x), PptInches(y), PptInches(4.05), PptInches(1.05))
        shape.fill.solid()
        shape.fill.fore_color.rgb = PptRGBColor.from_string(THEME["white"])
        shape.line.color.rgb = PptRGBColor.from_string(THEME["blue"])
        add_body_text(slide, [item], x + 0.18, y + 0.16, 3.7, 0.65, 12)

    # Slide 3
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PptRGBColor.from_string(THEME["white"])
    add_title(slide, "Pilot Metrics")
    chart_data = [
        ("Turnaround", "5", "2", "days"),
        ("Rework", "14", "4", "items"),
        ("Satisfaction", "3.4", "4.5", "score"),
        ("Volume", "18", "42", "docs"),
    ]
    for idx, (label, base, pilot, unit) in enumerate(chart_data):
        y = 1.05 + idx * 0.95
        add_body_text(slide, [label], 0.65, y + 0.11, 1.45, 0.35, 11)
        slide.shapes.add_shape(1, PptInches(2.25), PptInches(y + 0.18), PptInches(2.2), PptInches(0.22)).fill.solid()
        slide.shapes[-1].fill.fore_color.rgb = PptRGBColor.from_string("B8C2CC")
        width = 2.2 * (float(pilot) / max(float(base), float(pilot)))
        slide.shapes.add_shape(1, PptInches(4.75), PptInches(y + 0.18), PptInches(width), PptInches(0.22)).fill.solid()
        slide.shapes[-1].fill.fore_color.rgb = PptRGBColor.from_string(THEME["blue"])
        add_body_text(slide, [f"{base} -> {pilot} {unit}"], 7.25, y + 0.04, 1.95, 0.42, 12)

    # Slide 4
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PptRGBColor.from_string(THEME["light"])
    add_title(slide, "Risks and Controls")
    add_body_text(slide, ["Risks"], 0.8, 1.05, 3.4, 0.35, 16, THEME["blue"])
    add_body_text(slide, summary["risks"], 0.8, 1.55, 3.9, 2.0, 13)
    add_body_text(slide, ["Controls"], 5.25, 1.05, 3.4, 0.35, 16, THEME["blue"])
    add_body_text(slide, summary["actions"][1:], 5.25, 1.55, 3.9, 2.0, 13)
    slide.shapes.add_shape(1, PptInches(4.85), PptInches(1.05), PptInches(0.08), PptInches(3.35)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = PptRGBColor.from_string(THEME["cyan"])

    # Slide 5
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PptRGBColor.from_string(THEME["dark"])
    add_title(slide, "Recommendation", dark=True)
    add_body_text(slide, [summary["recommendation"]], 0.8, 1.2, 8.1, 1.1, 18, THEME["white"])
    for idx, action in enumerate(summary["actions"]):
        x = 0.85 + idx * 3.0
        slide.shapes.add_shape(9, PptInches(x), PptInches(3.1), PptInches(0.6), PptInches(0.6)).fill.solid()
        slide.shapes[-1].fill.fore_color.rgb = PptRGBColor.from_string(THEME["cyan"])
        add_body_text(slide, [str(idx + 1)], x + 0.2, 3.23, 0.2, 0.18, 14, THEME["dark"])
        add_body_text(slide, [action], x, 3.9, 2.35, 0.75, 11, THEME["white"])

    prs.save(PPTX_PATH)


def main():
    memo_text, metrics = read_inputs()
    summary = build_summary(memo_text, metrics)
    write_summary_md(summary)
    make_docx(summary, metrics)
    make_pdf(summary, metrics)
    make_pptx(summary, metrics)
    print(DOCX_PATH)
    print(PDF_PATH)
    print(PPTX_PATH)
    print(SUMMARY_MD_PATH)


if __name__ == "__main__":
    main()
